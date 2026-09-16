import io
import logging
import pandas as pd

logger = logging.getLogger(__name__)

def parse_tabular_or_docx(file_path_or_bytes, mime_type, file_name):
    """
    Parses Excel spreadsheets, CSV files, and Word documents.
    """
    try:
        lower_name = file_name.lower()
        if lower_name.endswith('.csv') or 'csv' in mime_type:
            df = pd.read_csv(file_path_or_bytes)
            text_repr = df.to_string(index=False)
            return {
                "success": True,
                "page_count": 1,
                "full_text": text_repr,
                "pages": [{"page_number": 1, "text": text_repr}]
            }
        
        elif lower_name.endswith(('.xlsx', '.xls')) or 'excel' in mime_type or 'spreadsheet' in mime_type:
            excel_file = pd.ExcelFile(file_path_or_bytes)
            all_text = []
            pages = []
            for idx, sheet_name in enumerate(excel_file.sheet_names):
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                sheet_text = f"Sheet: {sheet_name}\n" + df.to_string(index=False)
                all_text.append(sheet_text)
                pages.append({"page_number": idx + 1, "text": sheet_text})
            
            return {
                "success": True,
                "page_count": len(pages) or 1,
                "full_text": "\n\n".join(all_text),
                "pages": pages
            }

        elif lower_name.endswith(('.docx', '.doc')) or 'word' in mime_type:
            try:
                import docx
                doc = docx.Document(file_path_or_bytes)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                tables_text = []
                for table in doc.tables:
                    for row in table.rows:
                        tables_text.append(" | ".join(c.text.strip() for c in row.cells))
                
                full_text = "\n".join(paragraphs + tables_text)
                return {
                    "success": True,
                    "page_count": 1,
                    "full_text": full_text,
                    "pages": [{"page_number": 1, "text": full_text}]
                }
            except Exception as docx_err:
                logger.error(f"Error parsing Word doc: {docx_err}")
                return {
                    "success": False,
                    "page_count": 1,
                    "full_text": "",
                    "pages": [],
                    "error": str(docx_err)
                }

        # Plain text fallback
        if isinstance(file_path_or_bytes, bytes):
            text = file_path_or_bytes.decode('utf-8', errors='ignore')
        else:
            with open(file_path_or_bytes, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        return {
            "success": True,
            "page_count": 1,
            "full_text": text,
            "pages": [{"page_number": 1, "text": text}]
        }

    except Exception as e:
        logger.error(f"Error in tabular parser: {str(e)}")
        return {
            "success": False,
            "page_count": 1,
            "full_text": "",
            "pages": [],
            "error": str(e)
        }
